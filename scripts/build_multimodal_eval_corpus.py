#!/usr/bin/env python3
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "evals" / "multimodal_ingestion" / "corpus"


def _font(size: int):
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def build_corpus(output_dir: Path = OUTPUT_DIR) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest = {"documents": []}

    txt_path = output_dir / "benefits.txt"
    txt_path.write_text(
        "Employee benefit reference\nThe unique TXT policy code is TXT-731.\n",
        encoding="utf-8",
    )
    _record(manifest, txt_path, "TXT-731", "What is the unique TXT policy code?")

    md_path = output_dir / "operations.md"
    md_path.write_text(
        "# Operations\n\nThe unique Markdown runbook code is MD-204.\n\n"
        "## Recovery\n\nKeep the recovery steps in this section.\n",
        encoding="utf-8",
    )
    _record(manifest, md_path, "MD-204", "What is the unique Markdown runbook code?")

    json_path = output_dir / "customer.json"
    json_path.write_text(
        json.dumps(
            {
                "customer": {
                    "name": "Multimodal Test",
                    "verification_code": "JSON-442",
                    "status": "active",
                }
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    _record(manifest, json_path, "JSON-442", "What is the customer verification code?")

    docx_path = output_dir / "claims-guide.docx"
    document = Document()
    document.add_heading("Claims Guide", level=1)
    document.add_paragraph("The unique Word claims code is DOCX-431.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Escalation"
    table.cell(1, 1).text = "Manual review"
    document.save(docx_path)
    _normalize_zip_timestamps(docx_path)
    _record(manifest, docx_path, "DOCX-431", "What is the unique Word claims code?")

    image_path = output_dir / "invoice.png"
    image = Image.new("RGB", (1500, 520), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 100), "INVOICE VERIFICATION", font=_font(62), fill="black")
    draw.text((90, 220), "Unique image code: IMAGE-990", font=_font(54), fill="black")
    draw.text((90, 330), "Amount: 1280 TWD", font=_font(46), fill="black")
    image.save(image_path)
    _record(manifest, image_path, "IMAGE-990", "What is the unique image code?")

    pdf_path = output_dir / "financial-policy.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4, invariant=1)
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(72, 760, "Financial Policy")
    pdf.setFont("Helvetica", 14)
    pdf.drawString(72, 710, "The unique PDF policy code is PDF-620.")
    pdf.drawString(72, 680, "This page contains a searchable text layer.")
    pdf.save()
    _record(manifest, pdf_path, "PDF-620", "What is the unique PDF policy code?")

    scan_pdf_path = output_dir / "scanned-invoice.pdf"
    scan_pdf = canvas.Canvas(str(scan_pdf_path), pagesize=A4, invariant=1)
    scan_pdf.drawImage(
        ImageReader(str(image_path)),
        45,
        480,
        width=505,
        height=175,
        preserveAspectRatio=True,
        mask="auto",
    )
    scan_pdf.save()
    _record(
        manifest,
        scan_pdf_path,
        "IMAGE-990",
        "What is the unique code shown in the scanned invoice?",
        expected_method="pdf-ocr",
    )

    manifest_path = output_dir.parent / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return {
        "manifest": manifest_path.resolve().relative_to(PROJECT_ROOT).as_posix(),
        **manifest,
    }


def _record(manifest, path, marker, question, expected_method=""):
    try:
        portable_path = path.resolve().relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        portable_path = path.name
    manifest["documents"].append(
        {
            "path": portable_path,
            "filename": path.name,
            "marker": marker,
            "question": question,
            "expected_method": expected_method,
        }
    )


def _normalize_zip_timestamps(path: Path) -> None:
    with ZipFile(path, "r") as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]

    output = BytesIO()
    with ZipFile(output, "w") as target:
        for info, content in entries:
            normalized = ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            normalized.compress_type = info.compress_type
            normalized.external_attr = info.external_attr
            normalized.create_system = info.create_system
            target.writestr(normalized, content)
    path.write_bytes(output.getvalue())


if __name__ == "__main__":
    result = build_corpus()
    print(json.dumps(result, indent=2))
